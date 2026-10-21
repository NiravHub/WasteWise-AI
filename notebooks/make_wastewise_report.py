from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = "WasteWise_AI_Final_OEP_Report.docx"

doc = Document()

# ---------------------------------------------------------
# PAGE SETUP
# ---------------------------------------------------------
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(9.2)
doc.styles["Normal"].paragraph_format.space_after = Pt(3)

for name, size in [
    ("Title", 22),
    ("Heading 1", 14),
    ("Heading 2", 11),
]:
    doc.styles[name].font.name = "Aptos Display"
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True


# ---------------------------------------------------------
# HELPERS
# ---------------------------------------------------------
def shade(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells

        for col_index, value in enumerate(row):
            cells[col_index].text = str(value)
            cells[col_index].vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            if row_index == 0:
                shade(cells[col_index])

                for run in cells[col_index].paragraphs[0].runs:
                    run.bold = True

    return table


def heading(text, level=1):
    doc.add_heading(text, level=level)


def paragraph(text):
    doc.add_paragraph(text)


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def new_page():
    doc.add_page_break()


# =========================================================
# COVER PAGE
# =========================================================

for _ in range(2):
    paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SARVAJANIK COLLEGE OF ENGINEERING & TECHNOLOGY")
r.bold = True
r.font.size = Pt(15)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sarvajanik University, Surat")
r.bold = True
r.font.size = Pt(12)

paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("OPEN ENDED ASSIGNMENT IN ARTIFICIAL INTELLIGENCE")
r.bold = True
r.font.size = Pt(13)

paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("WasteWise AI")
r.bold = True
r.font.size = Pt(25)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Intelligent Multi-Class Waste Detection and Classification Using Neural Networks"
)
r.bold = True
r.font.size = Pt(15)

paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Computer Engineering Department")
r.bold = True
r.font.size = Pt(12)

paragraph("")

add_table([
    ["Team Member", "Contribution"],
    ["Drashtant Mevada", "Project development, ML pipeline and integration"],
    ["Ved Patel", "Team member"],
    ["Nirav Panwala", "Team member"],
    ["Raj Reshamwala", "Team member"],
])

paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Academic Year 2026–27").bold = True

new_page()


# =========================================================
# PAGE 2
# =========================================================

heading("1. Abstract")

paragraph(
    "WasteWise AI is an AI-powered waste intelligence system designed to detect "
    "waste objects, classify them into practical and detailed waste categories, "
    "reject unsupported inputs, and provide disposal guidance. The production "
    "pipeline combines YOLO11s object detection with an ImageNet-pretrained "
    "ResNet50 classifier, an embedding-based out-of-distribution (OOD) detector, "
    "and a semantic compatibility gate."
)

paragraph(
    "The detector uses six practical categories—plastic, glass, metal, "
    "paper_cardboard, biological, and other_trash—while the expanded ResNet50 "
    "classifier provides 19 detailed classes including medical-waste categories."
)

paragraph(
    "The project includes dataset integrity and duplicate-audit tooling, "
    "controlled image preprocessing, staged model training, OOD threshold "
    "calibration, external evaluation, and a FastAPI inference service. On the "
    "supported subset of the Medical Waste 4.0 external dataset, the Expanded "
    "ResNet50 Stage 2 achieved 99.14% classification accuracy, with 6,195 correct "
    "predictions out of 6,249 supported images."
)

paragraph(
    "The detector benchmark is substantially weaker than the classifier "
    "benchmark, so the external classification result must not be interpreted "
    "as end-to-end system accuracy."
)

heading("2. Introduction")

paragraph(
    "Waste segregation is difficult when images contain visually similar "
    "materials, mixed objects, medical waste, or items outside the training "
    "distribution. WasteWise AI was developed as a software-first computer-vision "
    "pipeline that detects objects, assigns a detailed waste class, evaluates "
    "whether the image belongs to the supported feature space, checks "
    "detector–classifier compatibility, and returns disposal guidance."
)

heading("3. Problem Statement")

paragraph(
    "Design and implement an AI-based waste analysis system capable of locating "
    "waste objects in an image, classifying them into multiple waste categories, "
    "identifying unsupported inputs, reducing inconsistent predictions, and "
    "translating accepted predictions into practical disposal guidance."
)

heading("4. Objectives")

bullets([
    "Build a multi-stage computer-vision pipeline for waste detection and classification.",
    "Use YOLO11s for object localization and ResNet50 for fine-grained classification.",
    "Support a 19-class detailed waste taxonomy including medical-waste classes.",
    "Implement calibrated OOD rejection using deep feature embeddings and cosine similarity.",
    "Add a semantic compatibility gate between detector and classifier outputs.",
    "Provide a deployable FastAPI backend for image inference and structured results.",
    "Evaluate the classifier on an external Medical Waste 4.0 dataset.",
])

new_page()


# =========================================================
# PAGE 3
# =========================================================

heading("5. Dataset and Data Preparation")

paragraph(
    "WasteWise AI uses multiple waste datasets and maintains dedicated audit "
    "tooling for dataset integrity, taxonomy mapping, annotation checks, duplicate "
    "provenance, and cross-split hash checks. The repository contains canonical "
    "manifests and validation outputs."
)

add_table([
    ["Component", "Purpose"],
    ["Canonical manifest", "Project-level image inventory and provenance."],
    ["Dataset integrity audits", "Check expected structures and image consistency."],
    ["Taxonomy audits", "Map source labels into WasteWise categories."],
    ["Duplicate/hash audits", "Check exact and candidate duplicate overlap."],
    ["Annotation audits", "Validate COCO/TACO-style detection annotations."],
])

heading("6. Medical Waste 4.0 Mapping", 2)

paragraph(
    "The Medical Waste 4.0 audit discovered 6,321 image files. Seven supported "
    "WasteWise categories account for 6,249 images; 72 images were excluded as "
    "unsupported."
)

add_table([
    ["Class", "Images"],
    ["gauze", "630"],
    ["medical_cap", "438"],
    ["medical_glasses", "486"],
    ["medical_gloves", "2,787"],
    ["shoe_cover", "972"],
    ["test_tube", "483"],
    ["urine_bag", "453"],
    ["Total supported", "6,249"],
    ["Unsupported", "72"],
])

heading("7. Image Preprocessing")

paragraph(
    "The controlled benchmark verifies the production preprocessing sequence: "
    "resize to 256, center-crop to 224×224, convert to tensor, and apply "
    "ImageNet normalization."
)

add_table([
    ["Step", "Configuration"],
    ["Resize", "256"],
    ["Center crop", "224 × 224"],
    ["Tensor conversion", "ToTensor"],
    ["Mean", "[0.485, 0.456, 0.406]"],
    ["Std", "[0.229, 0.224, 0.225]"],
])

heading("8. Class Taxonomies")

paragraph(
    "YOLO11s uses six practical classes: plastic, glass, metal, paper_cardboard, "
    "biological, and other_trash."
)

paragraph(
    "The Expanded ResNet50 uses 19 detailed classes: battery, biological, cable, "
    "cardboard, clothes, electronic_device, gauze, general_trash, glass, "
    "medical_cap, medical_glasses, medical_gloves, metal, paper, plastic, "
    "shoe_cover, shoes, test_tube, and urine_bag."
)

new_page()


# =========================================================
# PAGE 4
# =========================================================

heading("9. Proposed Methodology")

paragraph(
    "The production pipeline is: Image → YOLO11s detection → expanded object "
    "crop → Expanded ResNet50 classification → Expanded OOD detector → semantic "
    "compatibility gate → disposal guidance."
)

add_table([
    ["Stage", "Operation", "Output"],
    ["1", "Input image", "Image"],
    ["2", "YOLO11s", "Boxes + practical class + confidence"],
    ["3", "Crop expansion", "Classifier crop"],
    ["4", "ResNet50 Stage 2", "19-class prediction"],
    ["5", "OOD detector", "Support + similarity"],
    ["6", "Compatibility gate", "Accepted/rejected"],
    ["7", "Guidance", "Disposal recommendation"],
])

heading("10. YOLO11s Object Detection")

paragraph(
    "E4 detector validation contained 149 images and 407 instances. Precision "
    "was 0.216, recall 0.261, mAP@0.50 0.232, and mAP@0.50:0.95 0.188."
)

paragraph(
    "A TACO practical baseline reported precision 0.282, recall 0.387, "
    "mAP@0.50 0.306, and mAP@0.50:0.95 0.236."
)

heading("11. Expanded ResNet50 Classification")

paragraph(
    "The classifier uses an ImageNet-pretrained ResNet50 backbone with a "
    "replacement final classification layer. Stage 1 validation ran for "
    "25 epochs with batch size 32, learning rate 0.0005 and weight decay "
    "0.0001; the best validation macro-F1 was 95.75% at epoch 20."
)

paragraph(
    "Stage 2 checkpoint metadata records validation macro-F1 of 96.80% "
    "at epoch 6 over 19 classes."
)

heading("12. Out-of-Distribution Detection")

paragraph(
    "The OOD detector normalizes ResNet feature embeddings and compares them "
    "with known-class centroids using cosine similarity. The best similarity "
    "is compared with a calibrated threshold."
)

paragraph(
    "The expanded threshold is 0.230447, calibrated as the 1st percentile "
    "of correct-class cosine similarities on the expanded validation set."
)

paragraph(
    "A real rejection example produced similarity 0.1996 against threshold "
    "0.2304, so it was rejected because 0.1996 < 0.2304."
)

new_page()


# =========================================================
# PAGE 5
# =========================================================

heading("13. Semantic Compatibility Gate")

paragraph(
    "OOD rejection alone does not guarantee a semantically consistent result. "
    "WasteWise AI therefore checks whether the detailed ResNet prediction is "
    "compatible with the practical YOLO detection class."
)

paragraph(
    "The production predictor rejects an object when the OOD detector marks it "
    "unsupported or when the classifier output is incompatible with the detector "
    "class."
)

heading("14. Production Accounting")

paragraph(
    "The production accounting logic explicitly maintains:"
)

paragraph(
    "RAW DETECTIONS = ACCEPTED + OOD REJECTED + COMPATIBILITY REJECTED + OTHER REJECTED"
)

paragraph(
    "The detector confidence default is 0.35, while the expanded OOD calibration "
    "file supplies the OOD threshold."
)

heading("15. External Evaluation: Medical Waste 4.0")

paragraph(
    "The external evaluation discovered 6,321 images. After taxonomy filtering, "
    "6,249 supported images were evaluated and 72 were excluded. There were "
    "6,195 correct predictions, 54 incorrect predictions, and zero inference errors."
)

add_table([
    ["Measure", "Result"],
    ["Total discovered", "6,321"],
    ["Supported/evaluated", "6,249"],
    ["Unsupported/excluded", "72"],
    ["Correct", "6,195"],
    ["Incorrect", "54"],
    ["Inference errors", "0"],
    ["Accuracy on supported subset", "99.14%"],
])

heading("16. Per-Class External Results", 2)

add_table([
    ["Class", "Correct / Total", "Accuracy"],
    ["gauze", "600 / 630", "95.24%"],
    ["medical_cap", "438 / 438", "100%"],
    ["medical_glasses", "486 / 486", "100%"],
    ["medical_gloves", "2,764 / 2,787", "99.17%"],
    ["shoe_cover", "972 / 972", "100%"],
    ["test_tube", "482 / 483", "99.79%"],
    ["urine_bag", "453 / 453", "100%"],
])

paragraph(
    "Correct interpretation: on the supported subset of the Medical Waste 4.0 "
    "external dataset, the Expanded ResNet50 Stage 2 achieved 99.14% "
    "classification accuracy. This is not end-to-end detector accuracy."
)

new_page()


# =========================================================
# PAGE 6
# =========================================================

heading("17. Software Architecture and Implementation")

paragraph(
    "The system uses a React/TypeScript frontend and FastAPI inference backend. "
    "The backend provides GET /, GET /health, GET /ready and POST /predict."
)

paragraph(
    "It validates file extensions, empty files, image integrity and a 25 MB "
    "upload limit, initializes/reuses models, performs inference, and returns "
    "structured JSON."
)

add_table([
    ["Layer", "Technology"],
    ["Frontend", "React 19, TypeScript, Vite 8, Tailwind CSS 4"],
    ["UI", "lucide-react, motion, recharts, lenis"],
    ["Backend", "FastAPI"],
    ["Detector", "YOLO11s"],
    ["Classifier", "ImageNet-pretrained ResNet50"],
    ["OOD", "Centroids + cosine similarity"],
    ["API", "/, /health, /ready, /predict"],
])

heading("18. Production Inference Example")

paragraph(
    "For plastic_42.jpg, production inference produced 10 raw detections and "
    "2 accepted detections, with approximately 1,194.64 ms inference time."
)

paragraph(
    "Rejected examples included a detector-plastic result classified as "
    "general_trash with OOD similarity 0.3276 and another with no accepted "
    "classifier class and OOD similarity 0.1529."
)

heading("19. Disposal Guidance")

paragraph(
    "Accepted predictions are mapped to disposal categories and recommendations. "
    "The guidance layer operates downstream of model acceptance so unsupported "
    "or incompatible objects are not automatically converted into disposal advice."
)

heading("20. Key Project Artifacts")

bullets([
    "Model checkpoints for baseline and ResNet50 Stage 1/Stage 2 variants.",
    "Expanded OOD calibration configuration.",
    "Classification reports, confusion matrices and evaluation metrics.",
    "Canonical dataset manifest and integrity/audit outputs.",
    "COCO/TACO annotation audit outputs.",
    "Cross-dataset taxonomy and duplicate-provenance audits.",
    "Integrated evaluation and production inference outputs.",
    "FastAPI backend and React/TypeScript frontend.",
])

new_page()


# =========================================================
# PAGE 7
# =========================================================

heading("21. Advantages")

bullets([
    "Combines localization, classification and rejection logic.",
    "Supports six practical detector classes and 19 detailed classifier classes.",
    "Includes calibrated OOD rejection.",
    "Uses detector–classifier semantic compatibility checking.",
    "Maintains explicit accepted/rejected accounting.",
    "Includes external Medical Waste 4.0 evaluation.",
    "Provides a deployable API rather than only a notebook experiment.",
])

heading("22. Limitations")

bullets([
    "The current detector benchmark is comparatively weak, with mAP@0.50 of 0.232 on E4 validation.",
    "Classifier metrics and external classification accuracy do not represent end-to-end detection accuracy.",
    "External evaluation is reported only on the supported Medical Waste 4.0 subset.",
    "Mixed-waste scenes, occlusion, lighting variation and unseen waste types can challenge the system.",
    "Disposal guidance depends on the implemented taxonomy and rules.",
])

heading("23. Future Scope")

bullets([
    "Improve detector training and annotation coverage.",
    "Expand open-set evaluation with additional unrelated datasets.",
    "Improve calibration under domain shift.",
    "Add video inference and object tracking.",
    "Use active learning for uncertain/rejected samples.",
    "Evaluate latency and resource use on edge hardware.",
    "Expand disposal guidance using regional waste-management rules.",
])

heading("24. Conclusion")

paragraph(
    "WasteWise AI demonstrates a complete AI application pipeline for intelligent "
    "waste analysis. It integrates YOLO11s detection, Expanded ResNet50 "
    "classification, embedding-based OOD rejection, semantic compatibility "
    "checking, disposal guidance, external evaluation and API deployment."
)

paragraph(
    "The strongest verified result is 99.14% classification accuracy on the "
    "supported subset of Medical Waste 4.0. At the same time, the detector "
    "evaluation shows why the system should be described honestly as a multi-stage "
    "research prototype rather than a uniformly high-accuracy end-to-end detector."
)

heading("25. References / Evidence Base")

bullets([
    "WasteWise AI project source tree, model checkpoints, scripts, audit outputs and evaluation results.",
    "Medical Waste 4.0 integrity, structure, taxonomy mapping and external evaluation outputs.",
    "Controlled preprocessing and training-history/checkpoint metadata.",
    "YOLO detector validation and TACO practical baseline outputs.",
    "OOD calibration and production inference outputs.",
    "FastAPI backend and ResNet50 model implementation.",
])

paragraph(
    "Reporting note: historical metric values that could not be traced to a "
    "specific accessible project artifact were intentionally excluded. Numerical "
    "results in this report are limited to verified project evidence."
)


# ---------------------------------------------------------
# SAVE
# ---------------------------------------------------------
doc.save(OUTPUT)

print()
print("=" * 60)
print("REPORT CREATED SUCCESSFULLY")
print("=" * 60)
print(f"File: {OUTPUT}")
print()