from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pathlib import Path


# ============================================================
# WASTEWISE AI — PROFESSIONAL OEP PRESENTATION
# ============================================================

OUTPUT = "WasteWise_AI_Professional_OEP_Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]


# ============================================================
# COLOR SYSTEM
# ============================================================

BG = RGBColor(6, 12, 22)
BG_LIGHT = RGBColor(10, 20, 34)

WHITE = RGBColor(245, 248, 252)
TEXT = RGBColor(211, 222, 235)
MUTED = RGBColor(137, 157, 180)

CYAN = RGBColor(40, 220, 198)
GREEN = RGBColor(92, 226, 153)
BLUE = RGBColor(88, 151, 255)
AMBER = RGBColor(246, 190, 67)
RED = RGBColor(244, 91, 105)
PURPLE = RGBColor(166, 116, 255)

CARD = RGBColor(15, 28, 45)
CARD2 = RGBColor(18, 34, 54)
GRID = RGBColor(35, 54, 76)


# ============================================================
# BASIC HELPERS
# ============================================================

def add_background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

    # Top accent line
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        Inches(0.045)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = CYAN
    shape.line.fill.background()


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=WHITE,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font="Aptos"
):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    tf = box.text_frame
    tf.clear()

    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)

    tf.vertical_anchor = valign

    p = tf.paragraphs[0]
    p.alignment = align

    run = p.add_run()
    run.text = text

    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

    return box


def rounded_rect(
    slide,
    x,
    y,
    w,
    h,
    fill=CARD,
    line_color=None,
    radius=True
):
    shape_type = (
        MSO_SHAPE.ROUNDED_RECTANGLE
        if radius
        else MSO_SHAPE.RECTANGLE
    )

    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = fill

    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()

    return shape


def circle(
    slide,
    x,
    y,
    diameter,
    color
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

    return shape


def draw_line(
    slide,
    x1,
    y1,
    x2,
    y2,
    color=GRID,
    width=1
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2)
    )

    line.line.color.rgb = color
    line.line.width = Pt(width)

    return line


def arrow(
    slide,
    x1,
    y1,
    x2,
    y2,
    color=CYAN,
    width=2
):
    l = draw_line(
        slide,
        x1,
        y1,
        x2,
        y2,
        color,
        width
    )

    l.line.end_arrowhead = True

    return l


def pill(
    slide,
    text,
    x,
    y,
    w,
    color=CYAN
):
    rounded_rect(
        slide,
        x,
        y,
        w,
        0.32,
        color
    )

    add_text(
        slide,
        text,
        x + 0.04,
        y + 0.045,
        w - 0.08,
        0.20,
        8,
        BG,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE
    )


def section_header(
    slide,
    section,
    title,
    subtitle,
    number
):
    add_text(
        slide,
        section.upper(),
        0.65,
        0.38,
        3.5,
        0.22,
        9,
        CYAN,
        True
    )

    add_text(
        slide,
        title,
        0.65,
        0.77,
        11.3,
        0.60,
        28,
        WHITE,
        True
    )

    if subtitle:
        add_text(
            slide,
            subtitle,
            0.67,
            1.42,
            11.5,
            0.35,
            11,
            MUTED
        )

    add_text(
        slide,
        f"{number:02d}",
        12.0,
        0.38,
        0.55,
        0.22,
        9,
        MUTED,
        True,
        PP_ALIGN.RIGHT
    )


def footer(slide):
    draw_line(
        slide,
        0.65,
        7.08,
        12.65,
        7.08,
        GRID,
        0.7
    )

    add_text(
        slide,
        "WASTEWISE AI  •  SCET / SARVAJANIK UNIVERSITY",
        0.65,
        7.15,
        5.5,
        0.18,
        7,
        MUTED,
        True
    )

    add_text(
        slide,
        "OEP • 2026–27",
        10.7,
        7.15,
        1.95,
        0.18,
        7,
        MUTED,
        True,
        PP_ALIGN.RIGHT
    )


def metric_card(
    slide,
    x,
    y,
    w,
    h,
    value,
    label,
    accent,
    note=""
):
    rounded_rect(
        slide,
        x,
        y,
        w,
        h,
        CARD,
        GRID
    )

    rounded_rect(
        slide,
        x,
        y,
        0.055,
        h,
        accent,
        None,
        False
    )

    add_text(
        slide,
        value,
        x + 0.20,
        y + 0.16,
        w - 0.4,
        0.42,
        23,
        WHITE,
        True
    )

    add_text(
        slide,
        label.upper(),
        x + 0.20,
        y + 0.73,
        w - 0.4,
        0.20,
        8,
        accent,
        True
    )

    if note:
        add_text(
            slide,
            note,
            x + 0.20,
            y + 1.02,
            w - 0.4,
            0.28,
            7.8,
            MUTED
        )


def flow_box(
    slide,
    x,
    y,
    w,
    h,
    title,
    subtitle,
    accent
):
    rounded_rect(
        slide,
        x,
        y,
        w,
        h,
        CARD,
        GRID
    )

    circle(
        slide,
        x + 0.18,
        y + 0.17,
        0.19,
        accent
    )

    add_text(
        slide,
        title,
        x + 0.47,
        y + 0.10,
        w - 0.62,
        0.30,
        11,
        WHITE,
        True
    )

    add_text(
        slide,
        subtitle,
        x + 0.18,
        y + 0.52,
        w - 0.36,
        h - 0.60,
        8.5,
        MUTED
    )


def bullet(
    slide,
    text,
    x,
    y,
    w,
    accent=CYAN,
    size=10.5
):
    circle(
        slide,
        x,
        y + 0.10,
        0.07,
        accent
    )

    add_text(
        slide,
        text,
        x + 0.18,
        y,
        w - 0.18,
        0.32,
        size,
        TEXT
    )


# ============================================================
# SLIDE 1 — COVER
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

# Decorative grid
for i in range(8):
    draw_line(
        slide,
        8.15 + i * 0.55,
        1.05,
        8.15 + i * 0.55,
        6.35,
        RGBColor(17, 34, 52),
        0.45
    )

for i in range(9):
    draw_line(
        slide,
        8.15,
        1.05 + i * 0.59,
        12.65,
        1.05 + i * 0.59,
        RGBColor(17, 34, 52),
        0.45
    )


pill(
    slide,
    "OPEN ENDED ASSIGNMENT • ARTIFICIAL INTELLIGENCE",
    0.68,
    0.55,
    3.85,
    CYAN
)

add_text(
    slide,
    "WasteWise AI",
    0.68,
    1.32,
    6.9,
    0.70,
    44,
    WHITE,
    True
)

add_text(
    slide,
    "Intelligent Multi-Class Waste Detection\n"
    "and Classification Using Neural Networks",
    0.70,
    2.12,
    7.0,
    1.18,
    22,
    WHITE,
    True
)

add_text(
    slide,
    "A guarded computer-vision pipeline that detects waste, "
    "refines the class, rejects unsupported inputs, and returns disposal guidance.",
    0.72,
    3.50,
    6.75,
    0.75,
    12,
    MUTED
)

rounded_rect(
    slide,
    0.70,
    4.62,
    6.85,
    1.35,
    CARD,
    GRID
)

add_text(
    slide,
    "TEAM",
    0.96,
    4.84,
    0.70,
    0.20,
    8,
    CYAN,
    True
)

add_text(
    slide,
    "Drashtant Mevada  •  Smit Rathod  •  Nirav Panwala",
    0.96,
    5.15,
    5.95,
    0.23,
    10,
    WHITE,
    True
)

add_text(
    slide,
    "Krish Tandel  •  Ronak Solanki",
    0.96,
    5.48,
    5.95,
    0.23,
    10,
    WHITE,
    True
)

add_text(
    slide,
    "SCET • SARVAJANIK UNIVERSITY • 2026–27",
    0.72,
    6.72,
    5.9,
    0.22,
    8.5,
    MUTED,
    True
)

flow_box(
    slide,
    8.25,
    1.20,
    3.95,
    0.88,
    "YOLO11s",
    "6 practical detection classes",
    CYAN
)

arrow(
    slide,
    10.23,
    2.08,
    10.23,
    2.34
)

flow_box(
    slide,
    8.25,
    2.48,
    3.95,
    0.88,
    "Expanded ResNet50",
    "19 fine-grained classes",
    GREEN
)

arrow(
    slide,
    10.23,
    3.36,
    10.23,
    3.64
)

flow_box(
    slide,
    8.25,
    3.78,
    3.95,
    0.88,
    "Calibrated OOD",
    "Reject unsupported inputs",
    AMBER
)

arrow(
    slide,
    10.23,
    4.66,
    10.23,
    4.94
)

flow_box(
    slide,
    8.25,
    5.08,
    3.95,
    0.88,
    "Disposal Guidance",
    "Actionable accepted output",
    PURPLE
)


# ============================================================
# SLIDE 2 — PROBLEM
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "01 • Problem",
    "Waste recognition is a decision problem",
    "A useful system must know what it sees — and when not to trust what it sees.",
    2
)

rounded_rect(
    slide,
    0.68,
    2.05,
    7.20,
    4.45,
    CARD,
    GRID
)

add_text(
    slide,
    "THE GAP",
    0.96,
    2.34,
    1.0,
    0.22,
    8.5,
    CYAN,
    True
)

add_text(
    slide,
    "A single classifier can be confident\nfor the wrong reason.",
    0.96,
    2.72,
    5.95,
    0.75,
    20,
    WHITE,
    True
)

add_text(
    slide,
    "WasteWise AI therefore separates the problem into explicit stages.",
    0.96,
    3.62,
    5.95,
    0.35,
    11,
    MUTED
)

steps = [
    ("01", "Locate", "Find candidate objects in the image.", CYAN),
    ("02", "Refine", "Classify each crop into a detailed waste class.", GREEN),
    ("03", "Reject", "Reject unsupported or semantically incompatible predictions.", AMBER),
    ("04", "Guide", "Attach disposal guidance only after acceptance.", PURPLE),
]

yy = 4.15

for number, title, desc, accent in steps:

    circle(
        slide,
        0.98,
        yy,
        0.34,
        accent
    )

    add_text(
        slide,
        number,
        0.98,
        yy + 0.055,
        0.34,
        0.18,
        8,
        BG,
        True,
        PP_ALIGN.CENTER
    )

    add_text(
        slide,
        title,
        1.50,
        yy - 0.02,
        1.0,
        0.25,
        11.5,
        WHITE,
        True
    )

    add_text(
        slide,
        desc,
        2.35,
        yy - 0.02,
        4.9,
        0.34,
        9.7,
        MUTED
    )

    yy += 0.57


rounded_rect(
    slide,
    8.18,
    2.05,
    4.42,
    4.45,
    CARD2,
    GRID
)

add_text(
    slide,
    "DESIGN PRINCIPLE",
    8.48,
    2.34,
    2.1,
    0.22,
    8.5,
    AMBER,
    True
)

add_text(
    slide,
    "Don't turn uncertainty\ninto a fake label.",
    8.48,
    2.78,
    3.45,
    0.95,
    25,
    WHITE,
    True
)

draw_line(
    slide,
    8.48,
    3.95,
    12.12,
    3.95,
    GRID,
    0.8
)

pill(slide, "DETECTION", 8.48, 4.30, 2.0, CYAN)
pill(slide, "CLASSIFICATION", 8.48, 4.72, 2.0, GREEN)
pill(slide, "OOD REJECTION", 8.48, 5.14, 2.0, AMBER)
pill(slide, "COMPATIBILITY", 8.48, 5.56, 2.0, RED)

footer(slide)


# ============================================================
# SLIDE 3 — OBJECTIVES
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "02 • Objective",
    "One system. Four explicit objectives.",
    "The architecture is built around measurable stages rather than a black-box label.",
    3
)

objectives = [
    ("01", "DETECT", "Locate candidate waste objects with YOLO11s.", CYAN),
    ("02", "CLASSIFY", "Map object crops to 19 fine-grained classes with ResNet50.", GREEN),
    ("03", "REJECT", "Use calibrated OOD similarity and semantic compatibility.", AMBER),
    ("04", "GUIDE", "Return disposal guidance for accepted predictions.", PURPLE),
]

for i, (number, title, desc, accent) in enumerate(objectives):

    x = 0.68 + i * 3.13

    rounded_rect(
        slide,
        x,
        2.25,
        2.82,
        3.75,
        CARD,
        GRID
    )

    add_text(
        slide,
        number,
        x + 0.25,
        2.53,
        0.55,
        0.30,
        18,
        accent,
        True
    )

    add_text(
        slide,
        title,
        x + 0.25,
        3.12,
        2.2,
        0.35,
        14,
        WHITE,
        True
    )

    add_text(
        slide,
        desc,
        x + 0.25,
        3.72,
        2.25,
        1.05,
        11,
        TEXT
    )

    rounded_rect(
        slide,
        x + 0.25,
        5.42,
        2.25,
        0.06,
        accent,
        None,
        False
    )


add_text(
    slide,
    "CORE QUESTION",
    0.72,
    6.34,
    1.45,
    0.22,
    8.5,
    CYAN,
    True
)

add_text(
    slide,
    "Can the system make a useful prediction without silently overclaiming certainty?",
    2.10,
    6.30,
    9.8,
    0.30,
    11.5,
    WHITE,
    True
)

footer(slide)


# ============================================================
# SLIDE 4 — TAXONOMY
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "03 • Taxonomy",
    "Two vocabularies. Two jobs.",
    "The detector uses a practical object vocabulary; the classifier supplies fine-grained identity.",
    4
)

rounded_rect(
    slide,
    0.68,
    2.12,
    5.78,
    4.22,
    CARD,
    GRID
)

add_text(
    slide,
    "DETECTION TAXONOMY",
    0.96,
    2.42,
    2.4,
    0.22,
    8.5,
    CYAN,
    True
)

add_text(
    slide,
    "6 practical classes",
    0.96,
    2.78,
    3.0,
    0.4,
    22,
    WHITE,
    True
)

detector_classes = [
    ("plastic", CYAN),
    ("glass", GREEN),
    ("metal", AMBER),
    ("paper / cardboard", BLUE),
    ("biological", PURPLE),
    ("other trash", RED),
]

for i, (label, accent) in enumerate(detector_classes):

    col = i % 2
    row = i // 2

    pill(
        slide,
        label.upper(),
        0.96 + col * 2.55,
        3.48 + row * 0.68,
        2.20,
        accent
    )


rounded_rect(
    slide,
    6.72,
    2.12,
    5.88,
    4.22,
    CARD,
    GRID
)

add_text(
    slide,
    "CLASSIFICATION TAXONOMY",
    7.00,
    2.42,
    3.0,
    0.22,
    8.5,
    GREEN,
    True
)

add_text(
    slide,
    "19 fine-grained classes",
    7.00,
    2.78,
    3.6,
    0.4,
    22,
    WHITE,
    True
)

classifier_classes = [
    "battery",
    "biological",
    "cable",
    "cardboard",
    "clothes",
    "electronic device",
    "gauze",
    "general trash",
    "glass",
    "medical cap",
    "medical glasses",
    "medical gloves",
    "metal",
    "paper",
    "plastic",
    "shoe cover",
    "shoes",
    "test tube",
    "urine bag",
]

for i, label in enumerate(classifier_classes):

    col = i % 3
    row = i // 3

    add_text(
        slide,
        label,
        7.00 + col * 1.75,
        3.48 + row * 0.40,
        1.58,
        0.22,
        8.1,
        WHITE,
        True
    )


add_text(
    slide,
    "Detector = where / family",
    0.78,
    6.55,
    3.5,
    0.25,
    10.5,
    CYAN,
    True
)

add_text(
    slide,
    "Classifier = detailed identity",
    8.15,
    6.55,
    3.6,
    0.25,
    10.5,
    GREEN,
    True
)

footer(slide)


# ============================================================
# SLIDE 5 — PIPELINE
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "04 • Method",
    "The guarded inference pipeline",
    "Every accepted object passes through detection, classification, OOD screening and semantic compatibility.",
    5
)

nodes = [
    (0.68, "IMAGE", "Upload / capture", CYAN),
    (3.05, "YOLO11s", "Object detection", BLUE),
    (5.42, "CROP", "Expanded object crop", CYAN),
    (7.79, "RESNET50", "19-class prediction", GREEN),
    (10.16, "OOD + GATE", "Reject / accept", AMBER),
]

for x, title, subtitle, accent in nodes:

    flow_box(
        slide,
        x,
        2.55,
        1.95,
        1.32,
        title,
        subtitle,
        accent
    )


for x in [2.63, 5.00, 7.37, 9.74]:

    arrow(
        slide,
        x,
        3.21,
        x + 0.40,
        3.21
    )


rounded_rect(
    slide,
    3.00,
    4.48,
    7.25,
    1.12,
    CARD2,
    GRID
)

add_text(
    slide,
    "ACCEPTED OBJECT",
    3.28,
    4.74,
    1.55,
    0.22,
    8.5,
    GREEN,
    True
)

add_text(
    slide,
    "Class + detector confidence + classifier confidence + "
    "OOD similarity + guidance",
    4.72,
    4.70,
    5.15,
    0.42,
    10.5,
    WHITE,
    True
)

pill(
    slide,
    "OOD THRESHOLD 0.230447",
    4.35,
    5.85,
    2.65,
    AMBER
)

add_text(
    slide,
    "Calibration: 1st percentile of correct-class cosine similarities on the expanded validation set.",
    7.18,
    5.88,
    4.3,
    0.30,
    8.5,
    MUTED
)

footer(slide)


# ============================================================
# SLIDE 6 — DETECTOR RESULTS
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "05 • Detection",
    "The detector is the current system bottleneck",
    "E4 validation: 149 images and 407 annotated instances.",
    6
)

metric_card(
    slide,
    0.68,
    2.08,
    2.65,
    1.45,
    "21.6%",
    "Precision",
    CYAN
)

metric_card(
    slide,
    3.55,
    2.08,
    2.65,
    1.45,
    "26.1%",
    "Recall",
    GREEN
)

metric_card(
    slide,
    6.42,
    2.08,
    2.65,
    1.45,
    "23.2%",
    "mAP@0.50",
    AMBER
)

metric_card(
    slide,
    9.29,
    2.08,
    3.31,
    1.45,
    "18.8%",
    "mAP@0.50:0.95",
    PURPLE
)


rounded_rect(
    slide,
    0.68,
    3.82,
    7.35,
    2.50,
    CARD,
    GRID
)

add_text(
    slide,
    "PER-CLASS mAP@0.50",
    0.95,
    4.08,
    2.4,
    0.22,
    8.5,
    CYAN,
    True
)

detector_map = [
    ("plastic", 0.443, CYAN),
    ("glass", 0.016, RED),
    ("metal", 0.371, AMBER),
    ("paper / cardboard", 0.267, BLUE),
    ("other trash", 0.065, PURPLE),
]

yy = 4.53

for label, value, accent in detector_map:

    add_text(
        slide,
        label,
        0.96,
        yy - 0.02,
        1.45,
        0.22,
        8.7,
        WHITE,
        True
    )

    rounded_rect(
        slide,
        2.55,
        yy,
        4.65,
        0.20,
        BG_LIGHT,
        None,
        False
    )

    rounded_rect(
        slide,
        2.55,
        yy,
        4.65 * value,
        0.20,
        accent,
        None,
        False
    )

    add_text(
        slide,
        f"{value:.3f}",
        7.30,
        yy - 0.03,
        0.40,
        0.20,
        8.2,
        MUTED,
        True,
        PP_ALIGN.RIGHT
    )

    yy += 0.36


rounded_rect(
    slide,
    8.30,
    3.82,
    4.30,
    2.50,
    CARD2,
    GRID
)

add_text(
    slide,
    "READ THIS RESULT",
    8.58,
    4.08,
    2.0,
    0.22,
    8.5,
    AMBER,
    True
)

add_text(
    slide,
    "Strong classification\ncannot compensate for\nweak localization.",
    8.58,
    4.55,
    3.25,
    1.12,
    20,
    WHITE,
    True
)

add_text(
    slide,
    "Detector improvement is therefore the clearest next engineering target.",
    8.58,
    5.72,
    3.25,
    0.45,
    9.3,
    MUTED
)

footer(slide)


# ============================================================
# SLIDE 7 — RESNET
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "06 • Classification",
    "Expanded ResNet50 provides the fine-grained layer",
    "ImageNet-pretrained backbone with a task-specific 19-class output head.",
    7
)

rounded_rect(
    slide,
    0.68,
    2.10,
    4.15,
    4.28,
    CARD,
    GRID
)

add_text(
    slide,
    "MODEL",
    0.95,
    2.38,
    0.8,
    0.22,
    8.5,
    GREEN,
    True
)

add_text(
    slide,
    "ResNet50",
    0.95,
    2.78,
    2.7,
    0.45,
    27,
    WHITE,
    True
)

add_text(
    slide,
    "ImageNet-pretrained backbone\n"
    "Task-specific 19-class head\n"
    "Controlled image preprocessing",
    0.95,
    3.43,
    3.25,
    1.15,
    12,
    MUTED
)

metric_card(
    slide,
    0.95,
    5.02,
    1.55,
    1.05,
    "95.75%",
    "Stage 1 Macro-F1",
    GREEN
)

metric_card(
    slide,
    2.68,
    5.02,
    1.55,
    1.05,
    "96.80%",
    "Stage 2 Macro-F1",
    CYAN
)


rounded_rect(
    slide,
    5.08,
    2.10,
    7.52,
    4.28,
    CARD,
    GRID
)

add_text(
    slide,
    "PREPROCESSING",
    5.37,
    2.38,
    1.4,
    0.22,
    8.5,
    CYAN,
    True
)

preprocessing = [
    ("Resize\n256", CYAN),
    ("CenterCrop\n224", BLUE),
    ("Tensor", GREEN),
    ("ImageNet\nNormalize", PURPLE),
]

for i, (label, accent) in enumerate(preprocessing):

    x = 5.38 + i * 1.70

    rounded_rect(
        slide,
        x,
        2.93,
        1.42,
        0.90,
        BG_LIGHT,
        GRID
    )

    add_text(
        slide,
        label,
        x + 0.08,
        3.08,
        1.26,
        0.48,
        10,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE
    )

    if i < 3:

        arrow(
            slide,
            x + 1.42,
            3.38,
            x + 1.63,
            3.38
        )


add_text(
    slide,
    "NORMALIZATION CONSTANTS",
    5.37,
    4.30,
    2.2,
    0.22,
    8.5,
    AMBER,
    True
)

add_text(
    slide,
    "mean  [0.485, 0.456, 0.406]",
    5.37,
    4.72,
    3.0,
    0.25,
    10.2,
    WHITE,
    True
)

add_text(
    slide,
    "std   [0.229, 0.224, 0.225]",
    5.37,
    5.08,
    3.0,
    0.25,
    10.2,
    WHITE,
    True
)

add_text(
    slide,
    "Checkpoint evidence is reported as validation Macro-F1 — not test accuracy.",
    5.37,
    5.65,
    6.55,
    0.35,
    9.5,
    MUTED,
    True
)

footer(slide)


# ============================================================
# SLIDE 8 — OOD
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "07 • Safety",
    "OOD rejection asks: “Should we trust this label?”",
    "Known-class embeddings are compared with calibrated class centroids.",
    8
)

rounded_rect(
    slide,
    0.68,
    2.08,
    7.30,
    4.42,
    CARD,
    GRID
)

add_text(
    slide,
    "DECISION MECHANISM",
    0.96,
    2.36,
    2.0,
    0.22,
    8.5,
    AMBER,
    True
)

ood_steps = [
    ("Embedding", "Extract ResNet feature vector", CYAN),
    ("Normalize", "Convert to unit-length representation", BLUE),
    ("Compare", "Cosine similarity to each class centroid", GREEN),
    ("Threshold", "Best similarity must reach 0.230447", AMBER),
]

yy = 2.86

for title, subtitle, accent in ood_steps:

    flow_box(
        slide,
        0.98,
        yy,
        6.72,
        0.68,
        title,
        subtitle,
        accent
    )

    yy += 0.82


rounded_rect(
    slide,
    8.28,
    2.08,
    4.32,
    4.42,
    CARD2,
    GRID
)

add_text(
    slide,
    "REAL REJECTION",
    8.58,
    2.36,
    2.0,
    0.22,
    8.5,
    RED,
    True
)

add_text(
    slide,
    "Classifier output",
    8.58,
    2.82,
    1.6,
    0.22,
    8.7,
    MUTED
)

add_text(
    slide,
    "cardboard",
    10.35,
    2.79,
    1.65,
    0.28,
    13,
    WHITE,
    True
)

add_text(
    slide,
    "Best similarity",
    8.58,
    3.34,
    1.6,
    0.22,
    8.7,
    MUTED
)

add_text(
    slide,
    "0.1996",
    10.35,
    3.30,
    1.65,
    0.30,
    16,
    RED,
    True
)

add_text(
    slide,
    "Threshold",
    8.58,
    3.88,
    1.6,
    0.22,
    8.7,
    MUTED
)

add_text(
    slide,
    "0.2304",
    10.35,
    3.84,
    1.65,
    0.30,
    16,
    AMBER,
    True
)

draw_line(
    slide,
    8.58,
    4.40,
    12.15,
    4.40,
    GRID,
    0.8
)

add_text(
    slide,
    "0.1996 < 0.2304",
    8.58,
    4.76,
    3.1,
    0.38,
    20,
    RED,
    True
)

add_text(
    slide,
    "UNSUPPORTED IMAGE",
    8.58,
    5.35,
    3.1,
    0.35,
    14,
    WHITE,
    True
)

add_text(
    slide,
    "The model refuses to silently force the input into the known taxonomy.",
    8.58,
    5.84,
    3.30,
    0.42,
    9.5,
    MUTED
)

footer(slide)


# ============================================================
# SLIDE 9 — EXTERNAL VALIDATION
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "08 • External evaluation",
    "Medical Waste 4.0: 99.14% on the supported subset",
    "6,321 images were discovered; 6,249 mapped to supported WasteWise classes.",
    9
)

metric_card(
    slide,
    0.68,
    2.10,
    2.20,
    1.42,
    "6,321",
    "Discovered",
    CYAN
)

metric_card(
    slide,
    3.05,
    2.10,
    2.20,
    1.42,
    "6,249",
    "Supported",
    GREEN
)

metric_card(
    slide,
    5.42,
    2.10,
    2.20,
    1.42,
    "6,195",
    "Correct",
    BLUE
)

metric_card(
    slide,
    7.79,
    2.10,
    2.20,
    1.42,
    "54",
    "Incorrect",
    RED
)

metric_card(
    slide,
    10.16,
    2.10,
    2.44,
    1.42,
    "99.14%",
    "Accuracy",
    AMBER
)


rounded_rect(
    slide,
    0.68,
    3.82,
    7.30,
    2.50,
    CARD,
    GRID
)

add_text(
    slide,
    "SUPPORTED-SUBSET CLASS ACCURACY",
    0.95,
    4.08,
    3.2,
    0.22,
    8.5,
    CYAN,
    True
)

medical_classes = [
    ("gauze", 0.9524),
    ("medical cap", 1.0),
    ("medical glasses", 1.0),
    ("medical gloves", 0.9917),
    ("shoe cover", 1.0),
    ("test tube", 0.9979),
    ("urine bag", 1.0),
]

yy = 4.48

for label, value in medical_classes:

    add_text(
        slide,
        label,
        0.96,
        yy - 0.02,
        1.38,
        0.20,
        8.1,
        WHITE,
        True
    )

    rounded_rect(
        slide,
        2.45,
        yy,
        4.55,
        0.18,
        BG_LIGHT,
        None,
        False
    )

    rounded_rect(
        slide,
        2.45,
        yy,
        4.55 * value,
        0.18,
        GREEN,
        None,
        False
    )

    add_text(
        slide,
        f"{value * 100:.2f}%",
        7.18,
        yy - 0.03,
        0.45,
        0.20,
        7.9,
        MUTED,
        True,
        PP_ALIGN.RIGHT
    )

    yy += 0.27


rounded_rect(
    slide,
    8.28,
    3.82,
    4.32,
    2.50,
    CARD2,
    GRID
)

add_text(
    slide,
    "IMPORTANT WORDING",
    8.58,
    4.08,
    2.0,
    0.22,
    8.5,
    AMBER,
    True
)

add_text(
    slide,
    "99.14% is not\n“overall dataset accuracy.”",
    8.58,
    4.55,
    3.25,
    0.72,
    18,
    WHITE,
    True
)

add_text(
    slide,
    "It is classification accuracy on the supported subset of the external Medical Waste 4.0 benchmark.",
    8.58,
    5.38,
    3.25,
    0.65,
    10,
    MUTED
)

add_text(
    slide,
    "Inference errors: 0  •  Unsupported/excluded: 72",
    8.58,
    6.02,
    3.3,
    0.22,
    8.5,
    MUTED,
    True
)

footer(slide)


# ============================================================
# SLIDE 10 — DEPLOYMENT
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "09 • Deployment",
    "From image upload to an auditable API response",
    "FastAPI exposes the production pipeline through a single /predict endpoint.",
    10
)

architecture = [
    (0.68, "React UI", "Upload / visualize", CYAN),
    (3.00, "FastAPI", "POST /predict", BLUE),
    (5.32, "YOLO11s", "Detect", GREEN),
    (7.64, "ResNet50", "Classify", PURPLE),
    (9.96, "OOD + Gate", "Reject / accept", AMBER),
]

for x, title, subtitle, accent in architecture:

    flow_box(
        slide,
        x,
        2.28,
        1.95,
        1.15,
        title,
        subtitle,
        accent
    )

for x in [2.63, 4.95, 7.27, 9.59]:

    arrow(
        slide,
        x,
        2.86,
        x + 0.37,
        2.86
    )


rounded_rect(
    slide,
    0.68,
    4.00,
    5.95,
    2.25,
    CARD,
    GRID
)

add_text(
    slide,
    "PRODUCTION ACCOUNTING",
    0.96,
    4.28,
    2.7,
    0.22,
    8.5,
    CYAN,
    True
)

add_text(
    slide,
    "RAW DETECTIONS",
    0.96,
    4.75,
    1.4,
    0.22,
    9,
    MUTED,
    True
)

add_text(
    slide,
    "=",
    2.42,
    4.73,
    0.25,
    0.25,
    16,
    WHITE,
    True
)

pill(
    slide,
    "ACCEPTED",
    2.75,
    4.70,
    1.15,
    GREEN
)

add_text(
    slide,
    "+",
    4.05,
    4.73,
    0.25,
    0.25,
    16,
    WHITE,
    True
)

pill(
    slide,
    "OOD REJECTED",
    4.36,
    4.70,
    1.42,
    AMBER
)

add_text(
    slide,
    "+",
    5.91,
    4.73,
    0.25,
    0.25,
    16,
    WHITE,
    True
)

add_text(
    slide,
    "compatibility / other rejects",
    1.00,
    5.34,
    4.8,
    0.25,
    9.5,
    MUTED
)

add_text(
    slide,
    "Every raw detection is explicitly accounted for.",
    1.00,
    5.78,
    4.9,
    0.28,
    9.5,
    WHITE,
    True
)


rounded_rect(
    slide,
    6.92,
    4.00,
    5.68,
    2.25,
    CARD2,
    GRID
)

add_text(
    slide,
    "REAL INFERENCE SNAPSHOT",
    7.20,
    4.28,
    2.6,
    0.22,
    8.5,
    GREEN,
    True
)

metric_card(
    slide,
    7.20,
    4.73,
    1.45,
    1.03,
    "10",
    "Raw detections",
    CYAN
)

metric_card(
    slide,
    8.84,
    4.73,
    1.45,
    1.03,
    "2",
    "Accepted",
    GREEN
)

metric_card(
    slide,
    10.48,
    4.73,
    1.84,
    1.03,
    "1.19 s",
    "Inference time",
    AMBER
)

add_text(
    slide,
    "Example: plastic_42.jpg",
    7.20,
    5.94,
    3.0,
    0.22,
    8.7,
    MUTED,
    True
)

footer(slide)


# ============================================================
# SLIDE 11 — CONCLUSION
# ============================================================

slide = prs.slides.add_slide(BLANK)
add_background(slide)

section_header(
    slide,
    "10 • Conclusion",
    "The result is not just a label — it is a guarded decision",
    "WasteWise AI combines recognition, rejection and guidance in a deployable pipeline.",
    11
)


rounded_rect(
    slide,
    0.68,
    2.12,
    3.75,
    3.98,
    CARD,
    GRID
)

add_text(
    slide,
    "DELIVERED",
    0.96,
    2.42,
    1.3,
    0.22,
    8.5,
    GREEN,
    True
)

delivered = [
    "6-class object detection",
    "19-class fine classification",
    "Calibrated OOD rejection",
    "Semantic compatibility gate",
    "Disposal guidance",
    "FastAPI inference service",
]

for i, item in enumerate(delivered):

    bullet(
        slide,
        item,
        0.98,
        2.93 + i * 0.46,
        3.0,
        GREEN,
        10
    )


rounded_rect(
    slide,
    4.78,
    2.12,
    3.75,
    3.98,
    CARD,
    GRID
)

add_text(
    slide,
    "CURRENT LIMITATION",
    5.06,
    2.42,
    2.0,
    0.22,
    8.5,
    AMBER,
    True
)

add_text(
    slide,
    "Detection quality",
    5.06,
    2.95,
    2.6,
    0.35,
    20,
    WHITE,
    True
)

add_text(
    slide,
    "YOLO11s remains the main bottleneck in the full image-to-object pipeline.",
    5.06,
    3.55,
    2.95,
    0.85,
    11.5,
    MUTED
)

add_text(
    slide,
    "NEXT ENGINEERING PRIORITY",
    5.06,
    4.78,
    2.5,
    0.25,
    9,
    AMBER,
    True
)

add_text(
    slide,
    "Improve localization and small-object detection before claiming end-to-end robustness.",
    5.06,
    5.18,
    2.95,
    0.62,
    10.5,
    WHITE,
    True
)


rounded_rect(
    slide,
    8.88,
    2.12,
    3.72,
    3.98,
    CARD2,
    GRID
)

add_text(
    slide,
    "NEXT",
    9.16,
    2.42,
    0.7,
    0.22,
    8.5,
    PURPLE,
    True
)

future = [
    "Stronger detector",
    "Better small-object detection",
    "Larger balanced datasets",
    "More real-world validation",
    "Edge deployment",
]

for i, item in enumerate(future):

    bullet(
        slide,
        item,
        9.18,
        2.98 + i * 0.52,
        3.0,
        PURPLE,
        10.5
    )


add_text(
    slide,
    "“A useful AI system must know when not to guess.”",
    1.10,
    6.45,
    11.15,
    0.35,
    17,
    WHITE,
    True,
    PP_ALIGN.CENTER
)

add_text(
    slide,
    "THANK YOU",
    5.55,
    6.83,
    2.2,
    0.20,
    8.5,
    CYAN,
    True,
    PP_ALIGN.CENTER
)

footer(slide)


# ============================================================
# PRESENTATION METADATA
# ============================================================

prs.core_properties.title = (
    "WasteWise AI: Intelligent Multi-Class Waste Detection "
    "and Classification Using Neural Networks"
)

prs.core_properties.subject = (
    "Open Ended Assignment in Artificial Intelligence"
)

prs.core_properties.author = "WasteWise AI Team"

prs.core_properties.keywords = (
    "WasteWise AI, YOLO11s, ResNet50, OOD, "
    "waste classification, SCET"
)


# ============================================================
# SAVE
# ============================================================

prs.save(OUTPUT)

print()
print("=" * 70)
print("WASTEWISE AI PPT CREATED")
print("=" * 70)
print(f"File: {OUTPUT}")
print(f"Slides: {len(prs.slides)}")
print("Aspect ratio: 16:9")
print("=" * 70)